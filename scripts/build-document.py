#!/usr/bin/env python3
"""Build an editable DOCX from an ordered content plan, without rendering it."""
import argparse
import json
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True)
    parser.add_argument('--output', required=True)
    parser.add_argument('--font', default='Microsoft YaHei')
    args = parser.parse_args()
    output = Path(args.output).resolve()
    if output.exists():
        raise ValueError('Use a new document version; output already exists')
    data = json.loads(Path(args.input).read_text(encoding='utf-8'))
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    for name, size in [('Normal', 11.5), ('Title', 24), ('Subtitle', 12), ('Heading 1', 15), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = args.font
        style.font.size = Pt(size)
        style.font.bold = name in ('Title', 'Heading 1', 'Heading 2')
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + key), args.font)
        for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            fonts.attrib.pop(qn('w:' + key), None)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = Pt(18 if name == 'Normal' else size * 1.25)
        style.paragraph_format.keep_with_next = name != 'Normal'
        ppr = style.element.get_or_add_pPr()
        for border in ppr.findall(qn('w:pBdr')):
            ppr.remove(border)
    doc.add_paragraph(data['title'], 'Title')
    if data.get('subtitle'):
        doc.add_paragraph(data['subtitle'], 'Subtitle')
    for block in data['blocks']:
        kind = block['type']
        if kind == 'heading':
            doc.add_heading(block['text'], level=block.get('level', 1))
        elif kind in ('paragraph', 'metadata'):
            doc.add_paragraph(block['text'])
        elif kind == 'list':
            for index, item in enumerate(block['items'], 1):
                doc.add_paragraph(f'{index}. {item}' if block.get('ordered') else f'• {item}')
        elif kind == 'code':
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            run = paragraph.add_run(block['text'])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        elif kind == 'page_break':
            doc.add_page_break()
        else:
            raise ValueError(f'Unsupported document block: {kind}')
    doc.core_properties.title = data['title']
    doc.core_properties.author = ''
    doc.core_properties.comments = ''
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)
    print('DOCX built. Render and inspect every page before delivery.')


if __name__ == '__main__':
    try:
        main()
    except (ValueError, OSError, KeyError) as error:
        raise SystemExit(str(error))
