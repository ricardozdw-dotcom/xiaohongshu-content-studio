import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / 'scripts/workflow.py'
spec = importlib.util.spec_from_file_location('workflow', SCRIPT)
workflow = importlib.util.module_from_spec(spec)
spec.loader.exec_module(workflow)


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.job = Path(self.temp.name)
        for filename in ('source.docx', 'rewrite.docx'):
            doc = Document()
            doc.add_paragraph('匿名测试材料，不是真实用户任务。')
            doc.save(self.job / filename)
        self.run_action('init', '--rights', 'provided')

    def run_action(self, action, *args, ok=True):
        result = subprocess.run([sys.executable, str(SCRIPT), action, '--job', str(self.job), *args],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode == 0, ok, result.stdout + result.stderr)
        return result

    def source(self, coverage='complete'):
        self.run_action('submit-source', '--document', str(self.job / 'source.docx'), '--coverage', coverage)

    def ready(self):
        self.source()
        self.run_action('confirm-source', '--evidence', 'TEST FIXTURE ONLY: proceed to rewrite')
        self.run_action('submit-rewrite', '--document', str(self.job / 'rewrite.docx'))
        self.run_action('confirm-rewrite', '--evidence', 'TEST FIXTURE ONLY: approved for image production')
        return json.loads((self.job / 'job.json').read_text())

    def test_cannot_skip_source_review(self):
        self.run_action('submit-rewrite', '--document', str(self.job / 'rewrite.docx'), ok=False)

    def test_partial_capture_cannot_advance(self):
        self.source('partial')
        self.run_action('confirm-source', '--evidence', 'TEST FIXTURE ONLY', ok=False)

    def test_confirmation_needs_message(self):
        self.source()
        self.run_action('confirm-source', ok=False)

    def test_complete_review_and_hash_binding(self):
        state = self.ready()
        workflow.validate_ready(self.job, state)
        self.assertEqual(state['rewrite_approval']['sha256'], state['rewrite']['sha256'])

    def test_edit_after_approval_blocks_release(self):
        state = self.ready()
        doc = Document(self.job / 'rewrite.docx')
        doc.add_paragraph('Changed after review')
        doc.save(self.job / 'rewrite.docx')
        with self.assertRaises(ValueError):
            workflow.validate_ready(self.job, state)

    def test_resubmit_rewrite_resets_confirmation(self):
        self.ready()
        self.run_action('submit-rewrite', '--document', str(self.job / 'rewrite.docx'))
        state = json.loads((self.job / 'job.json').read_text())
        self.assertEqual(state['phase'], 'rewrite_review')
        self.assertNotIn('rewrite_approval', state)

    def test_resubmit_source_resets_both_confirmations(self):
        self.ready()
        self.source()
        state = json.loads((self.job / 'job.json').read_text())
        self.assertNotIn('source_approval', state)
        self.assertNotIn('rewrite_approval', state)

    def test_unknown_rights_disallows_verbatim(self):
        state = json.loads((self.job / 'job.json').read_text())
        state['rights'] = 'unknown'
        (self.job / 'job.json').write_text(json.dumps(state))
        self.run_action('submit-source', '--document', str(self.job / 'source.docx'),
                        '--coverage', 'complete', '--mode', 'verbatim', ok=False)

    def test_document_outside_job_is_rejected(self):
        with tempfile.TemporaryDirectory() as other:
            target = Path(other) / 'source.docx'
            target.write_bytes((self.job / 'source.docx').read_bytes())
            self.run_action('submit-source', '--document', str(target), '--coverage', 'complete', ok=False)

    def test_preview_cannot_be_packaged(self):
        self.ready()
        images = self.job / 'preview'
        images.mkdir()
        (images / 'render-report.json').write_text(json.dumps({'mode': 'preview', 'pages': []}))
        caption = self.job / '发布文案.md'
        caption.write_text('匿名测试文案', encoding='utf-8')
        result = subprocess.run([sys.executable, str(ROOT / 'scripts/package-release.py'),
                                 '--job', str(self.job), '--images', str(images),
                                 '--caption', str(caption), '--output', str(self.job / '发布包.zip')],
                                capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('Preview', result.stderr)

    def test_stale_approval_hash_is_rejected(self):
        state = self.ready()
        state['source_approval']['sha256'] = '0' * 64
        with self.assertRaises(ValueError):
            workflow.validate_ready(self.job, state)


class DocumentTests(unittest.TestCase):
    def test_technical_text_is_preserved_and_title_is_black(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            command = 'python --version\npython -c "print(\'你好\')"'
            plan = {'title': '教程步骤整理', 'blocks': [
                {'type': 'heading', 'text': '检查运行环境'},
                {'type': 'code', 'text': command},
                {'type': 'paragraph', 'text': '先核对版本，再执行示例。'}]}
            (root / 'plan.json').write_text(json.dumps(plan, ensure_ascii=False), encoding='utf-8')
            result = subprocess.run([sys.executable, str(ROOT / 'scripts/build-document.py'),
                                     '--input', str(root / 'plan.json'), '--output', str(root / 'result.docx')],
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            doc = Document(root / 'result.docx')
            self.assertIn(command, [p.text for p in doc.paragraphs])
            self.assertEqual(str(doc.styles['Title'].font.color.rgb), '000000')


if __name__ == '__main__':
    unittest.main()
